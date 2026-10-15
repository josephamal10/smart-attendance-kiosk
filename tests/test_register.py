"""Monthly register, Excel export, employee profile, backups."""
import io
import sqlite3
from datetime import date, datetime, timedelta

from models import AttendanceLog, Holiday, LeaveRequest
from services import backup

TODAY = date.today()
FIRST = TODAY.replace(day=1)


def _workday_before_today():
    d = TODAY - timedelta(days=1)
    while d.weekday() >= 5:
        d -= timedelta(days=1)
    return d


def _add(db_session, *objs):
    db = db_session()
    for o in objs:
        db.add(o)
    db.commit()
    db.close()


def _log(emp_id, day, t_in, t_out=None):
    return AttendanceLog(employee_id=emp_id, date=day.isoformat(), status="present",
                         in_time=datetime.combine(day, datetime.strptime(t_in, "%H:%M").time()),
                         out_time=datetime.combine(day, datetime.strptime(t_out, "%H:%M").time()) if t_out else None)


def test_register_cells_cover_every_status(client, make_employee, login, db_session):
    make_employee("boss", role="admin")
    a = make_employee("arun")
    day = _workday_before_today()
    if day.month != TODAY.month:  # early in the month: use today instead
        day = TODAY
    _add(db_session, _log(a.id, day, "09:40", "18:00"))
    data = client.get(f"/admin/register?year={TODAY.year}&month={TODAY.month}", headers=login("boss")).json()
    assert len(data["days"]) >= 28
    row = data["rows"][0]
    cell = row["cells"][day.day - 1]
    assert cell["status"] == "late" and cell["in"] == "09:40"
    statuses = {c["status"] for c in row["cells"]}
    assert statuses <= {"present", "late", "half_day", "absent", "leave", "holiday", "off", "future"}
    assert sum(row["totals"].values()) == len(data["days"])


def test_holiday_and_leave_beat_absent(client, make_employee, login, db_session):
    make_employee("boss", role="admin")
    a = make_employee("arun")
    _add(db_session,
         Holiday(date=FIRST, name="Founders day"),
         LeaveRequest(employee_id=a.id, start_date=FIRST + timedelta(days=1), end_date=FIRST + timedelta(days=1),
                      kind="sick", status="approved"))
    cells = client.get(f"/admin/register?year={TODAY.year}&month={TODAY.month}",
                       headers=login("boss")).json()["rows"][0]["cells"]
    assert cells[0]["status"] == "holiday" and cells[0]["name"] == "Founders day"
    if (FIRST + timedelta(days=1)).weekday() < 5:
        assert cells[1]["status"] == "leave"


def test_register_excel_opens_with_openpyxl(client, make_employee, login):
    from openpyxl import load_workbook
    make_employee("boss", role="admin")
    make_employee("arun", category="Developer")
    res = client.get("/admin/register.xlsx", headers=login("boss"))
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]
    ws = load_workbook(io.BytesIO(res.content)).active
    assert "attendance register" in ws["A1"].value
    assert ws["A5"].value == "Arun" and ws["B5"].value == "Developer"


def test_employee_profile(client, make_employee, login, db_session):
    make_employee("boss", role="admin")
    a = make_employee("arun", category="Tester")
    _add(db_session, _log(a.id, TODAY, "08:55"))
    p = client.get(f"/admin/employees/{a.id}/profile", headers=login("boss")).json()
    assert p["employee"]["name"] == "Arun" and p["employee"]["category"] == "Tester"
    assert p["row"]["cells"][TODAY.day - 1]["status"] == "present"
    assert p["face_photos"] == 0


def test_profile_of_admin_account_is_hidden(client, make_employee, login):
    boss = make_employee("boss", role="admin")
    assert client.get(f"/admin/employees/{boss.id}/profile", headers=login("boss")).status_code == 404


def test_backup_copies_a_readable_database(tmp_path, monkeypatch):
    src = tmp_path / "live.db"
    con = sqlite3.connect(src)
    con.execute("create table t(x)")
    con.execute("insert into t values (42)")
    con.commit()
    con.close()
    monkeypatch.setattr(backup, "BACKUP_DIR", str(tmp_path / "backups"))
    monkeypatch.setattr(backup, "KEEP", 2)
    names = [backup.backup_now(str(src)) for _ in range(1)]
    copy = tmp_path / "backups" / names[0]
    assert sqlite3.connect(copy).execute("select x from t").fetchone() == (42,)


def test_backups_are_pruned(tmp_path, monkeypatch):
    src = tmp_path / "live.db"
    sqlite3.connect(src).close()
    monkeypatch.setattr(backup, "BACKUP_DIR", str(tmp_path / "b"))
    monkeypatch.setattr(backup, "KEEP", 2)
    (tmp_path / "b").mkdir()
    for i in range(4):  # pretend older backups exist
        (tmp_path / "b" / f"live-2020-01-0{i + 1}_000000.db").write_bytes(b"")
    (tmp_path / "b" / "demo-2020-01-01_000000.db").write_bytes(b"")  # another database's backup
    backup.backup_now(str(src))
    assert len(list((tmp_path / "b").glob("live-*.db"))) == 2
    assert (tmp_path / "b" / "demo-2020-01-01_000000.db").exists()  # never touched
