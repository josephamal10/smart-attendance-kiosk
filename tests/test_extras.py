"""Work from home, one-day shift changes, birthdays & anniversaries, monthly register email."""
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

from auth import create_access_token
from models import AttendanceLog, Employee, LeaveRequest, ShiftChange
from routes.kiosk_routes import _same_day
from services import email_service, me_service, register, scheduler, settings_service
from services import rules as R

TODAY = date.today()


def _me(emp):
    return {"Authorization": "Bearer " + create_access_token({"sub": str(emp.id), "type": "self"})}


def _add(db_session, *objs):
    db = db_session()
    for o in objs:
        db.add(o)
    db.commit()
    db.close()


def _at(day, hhmm):
    return datetime.combine(day, datetime.strptime(hhmm, "%H:%M").time())


def _weekday_before_today(n=1):
    d = TODAY - timedelta(days=n)
    while d.weekday() >= 5:
        d -= timedelta(days=1)
    return d


# ── Work from home ───────────────────────────────────────────

def test_wfh_day_counts_as_attended_not_as_leave(client, make_employee, login, db_session):
    make_employee("boss", role="admin")
    arun = make_employee("arun")
    day = date(2026, 3, 4)  # a Wednesday
    _add(db_session, LeaveRequest(employee_id=arun.id, start_date=day, end_date=day, kind="wfh", status="approved"))
    row = client.get("/admin/register?year=2026&month=3", headers=login("boss")).json()["rows"][0]
    assert row["cells"][3]["status"] == "wfh" and row["totals"]["wfh"] == 1
    db = db_session()
    balance = me_service.leave_balance(db, db.get(Employee, arun.id), 2026)
    db.close()
    assert all(b["used"] == 0 for b in balance)            # doesn't use up leave allowance
    assert R.leave_days(db_session(), arun.id, day, day) == {}   # isn't leave


def test_wfh_request_from_the_dashboard(client, make_employee, login):
    make_employee("boss", role="admin")
    arun = make_employee("arun")
    res = client.post("/me/leave", headers=_me(arun), json={"start_date": TODAY.isoformat(),
                                                             "end_date": TODAY.isoformat(), "kind": "wfh"})
    assert res.status_code == 200 and res.json()["kind"] == "wfh"
    assert client.get("/me", headers=_me(arun)).json()["updates"][0]["title"].startswith("Work from home")


def test_wfh_shows_in_the_daily_email(db_session, make_employee):
    arun = make_employee("arun")
    _add(db_session, LeaveRequest(employee_id=arun.id, start_date=TODAY, end_date=TODAY, kind="wfh", status="approved"))
    db = db_session()
    s = email_service.build_summary(db, TODAY)
    db.close()
    assert s["rows"][0]["status"] == "Working from home" and s["counts"]["Working from home"] == 1


# ── One-day shift changes ────────────────────────────────────

def test_shift_change_flow_and_its_effect(client, make_employee, login, db_session):
    make_employee("boss", role="admin")
    make_employee("hr", role="viewer")
    arun = make_employee("arun", category="Tester")   # office hours 09:00–18:00
    day = _weekday_before_today()
    _add(db_session, AttendanceLog(employee_id=arun.id, date=day.isoformat(), status="present",
                                   in_time=_at(day, "11:05"), out_time=_at(day, "20:00")))
    req = client.post("/me/shift-change", headers=_me(arun),
                      json={"date": day.isoformat(), "start": "11:00", "end": "20:00", "reason": "Doctor"})
    assert req.status_code == 200 and req.json()["status"] == "pending"

    def flags():
        logs = client.get("/admin/attendance/all", headers=login("boss")).json()
        return logs[0]["flags"]
    assert "late" in flags()                         # still judged by 09:00 while pending

    sid = req.json()["id"]
    assert client.post(f"/admin/shift-changes/{sid}/decision", headers=login("hr"),
                       json={"status": "approved"}).status_code == 403
    ok = client.post(f"/admin/shift-changes/{sid}/decision", headers=login("boss"), json={"status": "approved"})
    assert ok.json()["status"] == "approved"
    assert flags() == []                              # on time for 11:00, not overtime till 20:00
    actions = [a["action"] for a in client.get("/admin/audit", headers=login("boss")).json()]
    assert "shift_change.approved" in actions


def test_shift_change_validation(client, make_employee):
    arun = make_employee("arun")
    bad = client.post("/me/shift-change", headers=_me(arun),
                      json={"date": TODAY.isoformat(), "start": "18:00", "end": "09:00"})
    assert bad.status_code == 422
    old = client.post("/me/shift-change", headers=_me(arun),
                      json={"date": (TODAY - timedelta(days=30)).isoformat(), "start": "10:00", "end": "19:00"})
    assert old.status_code == 422


def test_auto_checkout_uses_the_changed_end_time(db_session, make_employee):
    arun = make_employee("arun", category="Tester")
    day = _weekday_before_today()
    _add(db_session,
         AttendanceLog(employee_id=arun.id, date=day.isoformat(), status="present", in_time=_at(day, "11:00")),
         ShiftChange(employee_id=arun.id, date=day, start="11:00", end="20:00", status="approved",
                     decided_at=datetime.now()))
    db = db_session()
    assert scheduler.auto_checkout(db) == 1
    assert db.query(AttendanceLog).one().out_time == _at(day, "20:00")
    db.close()


# ── Birthdays & work anniversaries ───────────────────────────

def test_lock_screen_celebrations(client, make_employee, login, db_session):
    make_employee("boss", role="admin")
    arun = make_employee("arun kumar", category="Tester")
    mira = make_employee("mira", category="HR")
    db = db_session()
    db.get(Employee, arun.id).birthday = TODAY.replace(year=1998) if (TODAY.month, TODAY.day) != (2, 29) else date(2000, 2, 29)
    db.get(Employee, mira.id).joined_on = TODAY.replace(year=TODAY.year - 3) if (TODAY.month, TODAY.day) != (2, 29) else date(2024, 2, 29)
    db.commit()
    db.close()
    got = client.get("/kiosk/info").json()["celebrations"]
    assert {"kind": "birthday", "name": "Arun", "team": "Tester"} in got   # first name only
    assert any(c["kind"] == "anniversary" and c["name"] == "Mira" for c in got)

    s = client.get("/admin/settings", headers=login("boss")).json()
    s.pop("smtp_configured", None)
    client.put("/admin/settings", headers=login("boss"), json={**s, "show_celebrations": False})
    assert client.get("/kiosk/info").json()["celebrations"] == []


def test_leap_day_birthdays_are_celebrated_on_28_feb():
    leap_baby = date(2000, 2, 29)
    assert _same_day(leap_baby, date(2027, 2, 28))
    assert not _same_day(leap_baby, date(2028, 2, 28))   # 2028 has its own 29 Feb
    assert _same_day(leap_baby, date(2028, 2, 29))


def test_employee_birthday_and_join_date_can_be_set_and_cleared(client, make_employee, login):
    make_employee("boss", role="admin")
    h = login("boss")
    e = client.post("/admin/employees", headers=h, json={"name": "Divya", "category": "Tester",
                                                         "birthday": "1999-05-04", "joined_on": "2023-01-09"}).json()
    assert e["birthday"] == "1999-05-04" and e["joined_on"] == "2023-01-09"
    e = client.patch(f"/admin/employees/{e['id']}", headers=h, json={"birthday": None}).json()
    assert e["birthday"] is None and e["joined_on"] == "2023-01-09"


# ── Monthly register email ───────────────────────────────────

def test_monthly_register_is_emailed_once_on_the_1st(db_session, make_employee, monkeypatch, tmp_path):
    make_employee("arun")
    monkeypatch.setattr(email_service, "OUTBOX", str(tmp_path))
    monkeypatch.delenv("SMTP_USER", raising=False)
    db = db_session()
    settings_service.put(db, "email_enabled", True)
    db.commit()
    first = date(2026, 11, 1)
    early, later = _at(first, "08:00"), _at(first, "19:30")
    assert scheduler.maybe_send_monthly_register(db, early) is None          # before the email time
    assert "Saved to" in scheduler.maybe_send_monthly_register(db, later)
    assert scheduler.maybe_send_monthly_register(db, later + timedelta(days=3)) is None   # once a month
    ws = load_workbook(tmp_path / "attendance-2026-10.xlsx").active
    assert "October 2026" in ws["A1"].value and ws["A5"].value == "Arun"
    db.close()


def test_monthly_register_respects_its_switch(db_session, make_employee, monkeypatch, tmp_path):
    monkeypatch.setattr(email_service, "OUTBOX", str(tmp_path))
    db = db_session()
    settings_service.put(db, "email_enabled", True)
    settings_service.put(db, "monthly_register_email", False)
    db.commit()
    assert scheduler.maybe_send_monthly_register(db, _at(date(2026, 11, 1), "20:00")) is None
    db.close()


def test_register_excel_has_a_wfh_column():
    data = {"month_name": "March 2026", "work_start": "09:00", "work_end": "18:00",
            "days": [{"date": "2026-03-02", "day": 2, "weekday": "Mon", "working": True, "holiday": None}],
            "rows": [{"name": "Arun", "category": "Tester", "cells": [{"status": "wfh"}],
                      "totals": {s: (1 if s == "wfh" else 0) for s in register.STATUSES},
                      "attendance_pct": 100.0, "worked_hours": 0, "overtime_hours": 0}]}
    import io
    ws = load_workbook(io.BytesIO(register.to_xlsx(data, "NanoNino"))).active
    assert ws["C5"].value == "W" and ws.cell(row=4, column=7).value == "WFH"
